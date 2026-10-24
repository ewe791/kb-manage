"""知识单元导入：文件保存 -> 解析 -> 切片 -> 向量化 -> 入库。"""
from __future__ import annotations

import hashlib
import logging
import re
from datetime import datetime
from pathlib import Path
from uuid import uuid4

from docx import Document
from pypdf import PdfReader
from sqlalchemy.orm import Session

from app.config import settings
from app.models import KnowledgeChunk, KnowledgeUnit
from app.services.embedder import embed_texts, to_bytes
from app.services.vector_store import store

logger = logging.getLogger(__name__)

SUPPORTED_EXTS = {".pdf", ".md", ".markdown", ".docx", ".txt"}


def ext_of(filename: str) -> str:
    return Path(filename).suffix.lower()


def is_supported(filename: str) -> bool:
    return ext_of(filename) in SUPPORTED_EXTS


def parse_document(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
    if ext == ".docx":
        doc = Document(str(path))
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip()
    if ext in (".md", ".markdown", ".txt"):
        for encoding in ("utf-8", "utf-8-sig", "gb18030"):
            try:
                return path.read_text(encoding=encoding).strip()
            except UnicodeDecodeError:
                continue
        return path.read_text(encoding="utf-8", errors="ignore").strip()
    raise ValueError(f"不支持的文件格式：{ext}")


def split_paragraphs(text: str) -> list[str]:
    raw = text.replace("\r\n", "\n").replace("\r", "\n")
    blocks: list[str] = []
    for block in raw.split("\n"):
        block = block.strip()
        if not block:
            continue
        # 超长段落按句号切分，保证单段不超过切片上限
        if len(block) <= settings.CHUNK_SIZE:
            blocks.append(block)
            continue
        sentences = block.replace("！", "。").replace("？", "。").replace("；", "。").split("。")
        buf = ""
        for sentence in sentences:
            sentence = sentence.strip()
            if not sentence:
                continue
            sentence = sentence + "。"
            if len(buf) + len(sentence) > settings.CHUNK_SIZE and buf:
                blocks.append(buf)
                buf = sentence
            else:
                buf += sentence
        if buf:
            blocks.append(buf)
    return blocks


def chunk_text(text: str, size: int | None = None, overlap: int | None = None) -> list[str]:
    """按 500 字符 + 50 重叠切片，优先保留段落边界。"""
    size = settings.CHUNK_SIZE if size is None else size
    overlap = settings.CHUNK_OVERLAP if overlap is None else overlap
    paragraphs = split_paragraphs(text)
    if not paragraphs:
        return []

    chunks: list[str] = []
    current = ""
    for para in paragraphs:
        if not current:
            current = para
            continue
        if len(current) + len(para) + 1 <= size:
            current = f"{current}\n{para}"
        else:
            chunks.append(current)
            tail = current[-overlap:] if overlap > 0 else ""
            current = f"{tail}\n{para}".strip() if tail else para
    if current:
        chunks.append(current)

    # 仍超长的切片硬切（含重叠）
    result: list[str] = []
    for chunk in chunks:
        if len(chunk) <= size:
            result.append(chunk)
            continue
        start = 0
        while start < len(chunk):
            result.append(chunk[start : start + size])
            if start + size >= len(chunk):
                break
            start += size - overlap
    return [c.strip() for c in result if c.strip()]


def make_unit_code(title: str) -> str:
    stamp = datetime.now().strftime("%Y%m%d%H%M%S")
    digest = hashlib.md5(f"{title}{uuid4()}".encode("utf-8")).hexdigest()[:4].upper()
    return f"KU-{stamp}-{digest}"


_MD_NOISE = re.compile(r"[#>*`_~|\-]{1,}")


def _summary(text: str, limit: int = 200) -> str:
    """生成可读摘要：去掉 Markdown 标记与多余空白。"""
    clean = _MD_NOISE.sub(" ", text or "")
    clean = " ".join(clean.split())
    return clean[:limit]


def save_upload(user_id: str, filename: str, data: bytes) -> Path:
    user_dir = settings.upload_dir / user_id
    user_dir.mkdir(parents=True, exist_ok=True)
    safe_name = Path(filename).name
    target = user_dir / f"{uuid4().hex}_{safe_name}"
    target.write_bytes(data)
    return target


def _persist_chunks(db: Session, unit: KnowledgeUnit, content: str) -> int:
    """按 500/50 切片并写入向量，返回切片数。"""
    pieces = chunk_text(content)
    vectors = embed_texts(pieces) if pieces else []
    for idx, piece in enumerate(pieces):
        db.add(
            KnowledgeChunk(
                unit_id=unit.id,
                chunk_index=idx,
                chunk_text=piece,
                embedding=to_bytes(vectors[idx]) if len(vectors) > idx else None,
            )
        )
    return len(pieces)


def create_unit_from_text(
    db: Session,
    user_id: str,
    title: str,
    content: str,
    category: str = "未分类",
    summary: str | None = None,
    source_file_name: str | None = None,
    file_type: str = "md",
) -> dict:
    """不依赖文件的建单元入口（知识缺口一键建档等场景）。"""
    body = (content or "").strip()
    if not title.strip() or not body:
        raise ValueError("标题与正文不能为空")
    unit = KnowledgeUnit(
        unit_code=make_unit_code(title),
        title=title.strip()[:255],
        content=body,
        summary=(summary or _summary(body))[:512],
        category=category or "未分类",
        source_file_name=source_file_name,
        file_type=file_type,
        file_size=len(body.encode("utf-8")),
        status="active",
        creator_id=user_id,
    )
    db.add(unit)
    db.flush()
    chunk_count = _persist_chunks(db, unit, body)
    db.commit()
    db.refresh(unit)
    store.mark_dirty()
    return {"unit_id": unit.id, "title": unit.title, "chunk_count": chunk_count, "status": "success"}


def reindex_unit(db: Session, unit: KnowledgeUnit, content: str) -> int:
    """正文变更后重建切片与向量（保证检索同步）。"""
    body = (content or "").strip()
    if not body:
        raise ValueError("正文不能为空")
    unit.content = body
    unit.summary = _summary(body)
    unit.file_size = len(body.encode("utf-8"))
    db.query(KnowledgeChunk).filter(KnowledgeChunk.unit_id == unit.id).delete()
    count = _persist_chunks(db, unit, body)
    db.commit()
    db.refresh(unit)
    store.mark_dirty()
    return count


def import_bytes(
    db: Session,
    user_id: str,
    filename: str,
    data: bytes,
    category: str = "未分类",
    title: str | None = None,
) -> dict:
    """导入单个文件，返回 {file_name, unit_id, chunk_count, status}。"""
    safe_name = Path(filename).name
    if not is_supported(safe_name):
        return {"file_name": safe_name, "unit_id": None, "chunk_count": 0, "status": "unsupported"}
    if not data:
        return {"file_name": safe_name, "unit_id": None, "chunk_count": 0, "status": "empty"}

    path = save_upload(user_id, safe_name, data)
    try:
        content = parse_document(path)
    except Exception as exc:  # noqa: BLE001
        logger.exception("文档解析失败：%s", safe_name)
        return {
            "file_name": safe_name,
            "unit_id": None,
            "chunk_count": 0,
            "status": f"parse_failed: {exc}",
        }
    if not content.strip():
        return {"file_name": safe_name, "unit_id": None, "chunk_count": 0, "status": "parse_failed: 未解析到文本"}

    unit_title = title or Path(safe_name).stem
    unit = KnowledgeUnit(
        unit_code=make_unit_code(unit_title),
        title=unit_title[:255],
        content=content,
        summary=_summary(content),
        category=category or "未分类",
        source_file_name=safe_name,
        file_type=ext_of(safe_name).lstrip("."),
        file_size=len(data),
        status="active",
        creator_id=user_id,
    )
    db.add(unit)
    db.flush()

    pieces = _persist_chunks(db, unit, content)
    db.commit()
    db.refresh(unit)
    store.mark_dirty()
    return {
        "file_name": safe_name,
        "unit_id": unit.id,
        "chunk_count": pieces,
        "status": "success",
        "title": unit.title,
    }


def import_path(db: Session, user_id: str, file_path: str, category: str = "未分类") -> dict:
    path = Path(file_path)
    if not path.is_absolute():
        path = settings.data_dir.parent / path
    if not path.exists():
        return {"file_name": file_path, "unit_id": None, "chunk_count": 0, "status": "file_not_found"}
    return import_bytes(db, user_id, path.name, path.read_bytes(), category=category)


def delete_units(db: Session, unit_ids: list[str]) -> int:
    if not unit_ids:
        return 0
    units = db.query(KnowledgeUnit).filter(KnowledgeUnit.id.in_(unit_ids)).all()
    for unit in units:
        db.delete(unit)
    db.commit()
    store.mark_dirty()
    return len(units)
